"""POST /file-processing/process: selective document/file extraction, ported from ai-platform's
app/rag/file_router.py + ingest.py (the FILE TYPE ROUTER). The core idea, straight from that
reference: use deterministic local readers for everything a library can already do reliably, and
reserve Qwen Omni for the parts that genuinely need visual understanding — so Omni cost and
latency scale with how much of the file actually needs vision, not with total page count.

  pdf     PyMuPDF text layer per page. A page is rendered whole and OCR'd only when its text
          layer is (near) empty OR it carries one embedded image covering a large share of the
          page (LARGE_IMAGE_PAGE_FRACTION) — the second case catches a scanned page sitting under
          a short native caption, which text-length alone would miss (ai-platform hit this for
          real on a document shaped exactly that way). A page that keeps its text layer still
          gets its MEDIUM-sized embedded images (bigger than a bullet glyph, smaller than
          "dominates the page") pulled out and OCR'd individually, so a chart next to body text
          isn't silently dropped either way.
  office  xlsx/docx/pptx: native readers (openpyxl/python-docx/python-pptx) for the text, PLUS
          Omni vision analysis of every embedded image big enough to be real content rather than
          a decorative icon or divider.
  image / video / audio   Sent to Omni whole — no local reader can do this one.
  text    Local UTF-8 decode, no Omni involved at all.

Rendering (PyMuPDF rasterizing a page to PNG) costs no Omni tokens, but does cost CPU/memory for
that render — real, just not billed per page the way an Omni call is.

Unlike ai-platform's ingest.py, this has no chunking, embedding, or database write (there's
nothing to index here — it's a playground page, not a RAG pipeline) and tracks every individual
Omni call's own latency/token usage, surfaced per page/item plus an aggregate, so the key metric
the reference recommends benchmarking — the fraction of a document that actually needed Omni, not
the total page count — is directly visible in the response."""
import asyncio
import io
import time
from dataclasses import dataclass, field
from pathlib import Path

import pymupdf
from docx import Document as DocxDocument
from fastapi import APIRouter, File, HTTPException, UploadFile
from openpyxl import load_workbook
from PIL import Image as PILImage
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

import omni_video

router = APIRouter()

IMAGE_EXTS = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
              ".webp": "image/webp", ".gif": "image/gif", ".bmp": "image/bmp"}
AUDIO_EXTS = {".mp3": "mp3", ".wav": "wav", ".m4a": "m4a", ".aac": "aac", ".ogg": "ogg", ".flac": "flac"}
VIDEO_EXTS = {".mp4": "video/mp4", ".mov": "video/quicktime", ".webm": "video/webm", ".mkv": "video/x-matroska"}
OFFICE_EXTS = (".xlsx", ".xlsm", ".docx", ".pptx")

# Same defaults as ai-platform's ingest.py (see its own comments for the reasoning behind each).
MIN_TEXT_LAYER_CHARS = 20
LARGE_IMAGE_PAGE_FRACTION = 0.3
MIN_EMBEDDED_IMAGE_FRACTION = 0.02
MIN_EMBEDDED_IMAGE_PIXELS = 2500
OCR_RENDER_SCALE = 2.0  # 144 DPI — ai-platform's tuned value; 200 DPI is a reasonable starting
# point in general, but this is the reference's own verified setting, not a guess.
MAX_OCR_PAGES = 30
MAX_EMBEDDED_IMAGES = 30
OCR_CONCURRENCY = 4


def _ext(filename: str | None) -> str:
    name = (filename or "").lower()
    dot = name.rfind(".")
    return name[dot:] if dot != -1 else ""


def route_file(filename: str | None, data: bytes) -> str:
    """One of 'pdf' | 'image' | 'video' | 'audio' | 'office' | 'text'."""
    ext = _ext(filename)
    if ext == ".pdf" or data[:5] == b"%PDF-":
        return "pdf"
    if ext in IMAGE_EXTS:
        return "image"
    if ext in VIDEO_EXTS:
        return "video"
    if ext in AUDIO_EXTS:
        return "audio"
    if ext in OFFICE_EXTS:
        return "office"
    return "text"


@dataclass
class OmniCall:
    """One Omni call's own stats, tracked regardless of what triggered it (OCR'd page, embedded
    image, standalone image/audio/video) — this is what lets the response show the reference's
    recommended benchmark metric directly: how much of the document actually needed Omni."""
    kind: str  # "page" | "image" | "audio" | "video"
    label: str  # e.g. "page 3", "image 2 on slide 4"
    latency_ms: int | None
    tokens: dict | None
    error: str | None


@dataclass
class Unit:
    """One page / sheet / slide / whole-file result."""
    index: int | None
    content_type: str  # "text" | "ocr" | "image_caption" | "transcript"
    text: str
    omni_used: bool = False


@dataclass
class ProcessResult:
    kind: str
    units: list[Unit] = field(default_factory=list)
    omni_calls: list[OmniCall] = field(default_factory=list)


def _has_large_embedded_image(page) -> bool:
    try:
        page_area = page.rect.width * page.rect.height
        if page_area <= 0:
            return False
        for info in page.get_image_info():
            bbox = info.get("bbox")
            if not bbox:
                continue
            x0, y0, x1, y1 = bbox
            area = max(0.0, x1 - x0) * max(0.0, y1 - y0)
            if area / page_area >= LARGE_IMAGE_PAGE_FRACTION:
                return True
    except Exception:
        pass  # malformed page image info must never fail the whole extraction
    return False


def _medium_embedded_images(doc, page) -> list[tuple[bytes, str]]:
    out: list[tuple[bytes, str]] = []
    try:
        page_area = page.rect.width * page.rect.height
        if page_area <= 0:
            return out
        seen_xrefs: set[int] = set()
        for info in page.get_image_info(xrefs=True):
            xref = info.get("xref")
            bbox = info.get("bbox")
            if not xref or not bbox or xref in seen_xrefs:
                continue
            x0, y0, x1, y1 = bbox
            fraction = (max(0.0, x1 - x0) * max(0.0, y1 - y0)) / page_area
            if MIN_EMBEDDED_IMAGE_FRACTION <= fraction < LARGE_IMAGE_PAGE_FRACTION:
                seen_xrefs.add(xref)
                extracted = doc.extract_image(xref)
                out.append((extracted["image"], extracted["ext"]))
    except Exception:
        pass
    return out


def _is_trivial_image(blob: bytes) -> bool:
    """True if almost certainly a decorative icon/bullet/divider rather than real content.
    Unreadable bytes count as NOT trivial — better to waste one Omni call than silently drop
    something PIL merely failed to parse."""
    try:
        with PILImage.open(io.BytesIO(blob)) as im:
            return (im.width * im.height) < MIN_EMBEDDED_IMAGE_PIXELS
    except Exception:
        return False


def _image_mime(blob: bytes) -> str:
    try:
        with PILImage.open(io.BytesIO(blob)) as im:
            fmt = (im.format or "PNG").lower()
        return f"image/{'jpeg' if fmt == 'jpg' else fmt}"
    except Exception:
        return "image/png"


def _join_text_and_notes(text: str, notes: list[str]) -> tuple[str, str]:
    """(joined_text, content_type) — notes stand alone (content_type becomes 'image_caption')
    when there's no native text at all, e.g. a slide that's nothing but a picture."""
    if not notes:
        return text, "text"
    body = "\n\n".join(notes)
    return (f"{text}\n\n{body}" if text else body), ("image_caption" if not text else "text")


def _read_pdf(data: bytes) -> list[tuple[int, str, bytes | None, list[tuple[bytes, str]]]]:
    """(page_number, text_layer, full_page_png_or_None, medium_embedded_images) per page."""
    pages: list[tuple[int, str, bytes | None, list[tuple[bytes, str]]]] = []
    with pymupdf.open(stream=data, filetype="pdf") as doc:
        for i, page in enumerate(doc, start=1):
            try:
                text = (page.get_text("text") or "").strip()
            except Exception:
                text = ""
            png = None
            medium_images: list[tuple[bytes, str]] = []
            if len(text) < MIN_TEXT_LAYER_CHARS or _has_large_embedded_image(page):
                try:
                    png = page.get_pixmap(matrix=pymupdf.Matrix(OCR_RENDER_SCALE, OCR_RENDER_SCALE)).tobytes("png")
                except Exception:
                    png = None
            else:
                medium_images = _medium_embedded_images(doc, page)
            pages.append((i, text, png, medium_images))
    return pages


async def _ocr_page(sem: asyncio.Semaphore, calls: list[OmniCall], number: int, png: bytes) -> str:
    async with sem:
        try:
            text, usage, latency_ms = await omni_video.transcribe_page(png, mime="image/png")
            calls.append(OmniCall("page", f"page {number}", latency_ms, usage, None))
            return text
        except omni_video.OmniError as err:
            calls.append(OmniCall("page", f"page {number}", None, None, str(err)))
            return ""


async def _analyze_embedded(sem: asyncio.Semaphore, calls: list[OmniCall], label: str, blob: bytes, mime: str) -> str:
    async with sem:
        try:
            text, usage, latency_ms = await omni_video.analyze_image(blob, mime=mime)
            calls.append(OmniCall("image", label, latency_ms, usage, None))
            return text
        except omni_video.OmniError as err:
            calls.append(OmniCall("image", label, None, None, str(err)))
            return ""


async def _extract_pdf(data: bytes, calls: list[OmniCall]) -> list[Unit]:
    pdf_pages = await asyncio.to_thread(_read_pdf, data)
    sem = asyncio.Semaphore(OCR_CONCURRENCY)

    scanned = [(n, png) for n, _t, png, _imgs in pdf_pages if png][:MAX_OCR_PAGES]
    ocr_texts = await asyncio.gather(*(_ocr_page(sem, calls, n, png) for n, png in scanned))
    ocr = {n: t for (n, _png), t in zip(scanned, ocr_texts)}

    img_jobs = [
        (n, j, img, f"image/{'jpeg' if ext in ('jpg', 'jpeg') else ext}")
        for n, _t, png, imgs in pdf_pages if not png
        for j, (img, ext) in enumerate(imgs, start=1)
    ][:MAX_EMBEDDED_IMAGES]
    analyzed = await asyncio.gather(*(
        _analyze_embedded(sem, calls, f"image {j} on page {n}", img, mime) for n, j, img, mime in img_jobs
    ))
    image_notes: dict[int, list[str]] = {}
    for (n, j, _img, _mime), analysis in zip(img_jobs, analyzed):
        if analysis.strip():
            image_notes.setdefault(n, []).append(f"[Image {j} on page {n}]\n{analysis}")

    units: list[Unit] = []
    for n, text, png, _imgs in pdf_pages:
        if png and ocr.get(n, "").strip():
            units.append(Unit(n, "ocr", ocr[n], omni_used=True))
        else:
            joined, ctype = _join_text_and_notes(text, image_notes.get(n, []))
            units.append(Unit(n, ctype, joined, omni_used=bool(image_notes.get(n))))
    return units


async def _extract_pptx(data: bytes, calls: list[OmniCall]) -> list[Unit]:
    prs = Presentation(io.BytesIO(data))
    slide_area = prs.slide_width * prs.slide_height
    sem = asyncio.Semaphore(OCR_CONCURRENCY)
    units: list[Unit] = []

    for i, slide in enumerate(prs.slides, start=1):
        try:
            lines = [
                shape.text_frame.text for shape in slide.shapes
                if shape.has_text_frame and shape.text_frame.text.strip()
            ]
            text = "\n".join(lines)

            image_jobs: list[tuple[bytes, str]] = []
            for shape in slide.shapes:
                if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                    continue
                if slide_area > 0 and (shape.width * shape.height) / slide_area < MIN_EMBEDDED_IMAGE_FRACTION:
                    continue
                try:
                    image_jobs.append((shape.image.blob, f"image/{shape.image.ext}"))
                except Exception:
                    pass
            image_jobs = image_jobs[:MAX_EMBEDDED_IMAGES]

            analyzed = await asyncio.gather(*(
                _analyze_embedded(sem, calls, f"image {j} on slide {i}", blob, mime)
                for j, (blob, mime) in enumerate(image_jobs, start=1)
            ))
            notes = [f"[Image {j} on slide {i}]\n{a}" for j, a in enumerate(analyzed, start=1) if a.strip()]
            joined, ctype = _join_text_and_notes(text, notes)
            units.append(Unit(i, ctype, joined, omni_used=bool(notes)))
        except Exception:
            units.append(Unit(i, "text", "", omni_used=False))
    return units


async def _extract_docx(data: bytes, calls: list[OmniCall]) -> list[Unit]:
    doc = DocxDocument(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())

    image_jobs: list[tuple[bytes, str]] = []
    for rel in doc.part.rels.values():
        if rel.is_external or "image" not in rel.reltype:
            continue
        try:
            blob = rel.target_part.blob
        except Exception:
            continue
        if not _is_trivial_image(blob):
            image_jobs.append((blob, rel.target_part.content_type))
    image_jobs = image_jobs[:MAX_EMBEDDED_IMAGES]

    sem = asyncio.Semaphore(OCR_CONCURRENCY)
    analyzed = await asyncio.gather(*(
        _analyze_embedded(sem, calls, f"image {j}", blob, mime)
        for j, (blob, mime) in enumerate(image_jobs, start=1)
    ))
    notes = [f"[Image {j}]\n{a}" for j, a in enumerate(analyzed, start=1) if a.strip()]
    joined, ctype = _join_text_and_notes(text, notes)
    return [Unit(None, ctype, joined, omni_used=bool(notes))]


async def _extract_xlsx(data: bytes, calls: list[OmniCall]) -> list[Unit]:
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    sem = asyncio.Semaphore(OCR_CONCURRENCY)
    units: list[Unit] = []

    for i, ws in enumerate(wb.worksheets, start=1):
        try:
            lines = []
            for row in ws.iter_rows(values_only=True):
                cells = ["" if v is None else str(v) for v in row]
                while cells and not cells[-1]:
                    cells.pop()
                if cells:
                    lines.append("\t".join(cells))
            text = (f"Sheet: {ws.title}\n" + "\n".join(lines)) if lines else ""

            image_jobs: list[bytes] = []
            for im in getattr(ws, "_images", []):
                try:
                    blob = im._data()
                except Exception:
                    continue
                if not _is_trivial_image(blob):
                    image_jobs.append(blob)
            image_jobs = image_jobs[:MAX_EMBEDDED_IMAGES]

            analyzed = await asyncio.gather(*(
                _analyze_embedded(sem, calls, f"image {j} on sheet {ws.title}", blob, _image_mime(blob))
                for j, blob in enumerate(image_jobs, start=1)
            ))
            notes = [f"[Image {j} on sheet {ws.title}]\n{a}" for j, a in enumerate(analyzed, start=1) if a.strip()]
            joined, ctype = _join_text_and_notes(text, notes)
            units.append(Unit(i, ctype, joined, omni_used=bool(notes)))
        except Exception:
            units.append(Unit(i, "text", "", omni_used=False))
    return units


async def process_file(data: bytes, *, filename: str | None) -> ProcessResult:
    kind = route_file(filename, data)
    calls: list[OmniCall] = []
    ext = _ext(filename)

    if kind == "pdf":
        units = await _extract_pdf(data, calls)
    elif kind == "office":
        if ext in (".xlsx", ".xlsm"):
            units = await _extract_xlsx(data, calls)
        elif ext == ".docx":
            units = await _extract_docx(data, calls)
        elif ext == ".pptx":
            units = await _extract_pptx(data, calls)
        else:
            units = []
    elif kind == "image":
        try:
            text, usage, latency_ms = await omni_video.analyze_image(data, mime=IMAGE_EXTS.get(ext, "image/png"))
            calls.append(OmniCall("image", filename or "image", latency_ms, usage, None))
            units = [Unit(None, "image_caption", text, omni_used=True)]
        except omni_video.OmniError as err:
            calls.append(OmniCall("image", filename or "image", None, None, str(err)))
            units = [Unit(None, "image_caption", "", omni_used=True)]
    elif kind == "audio":
        try:
            text, usage, latency_ms = await omni_video.transcribe_audio(data, audio_format=AUDIO_EXTS.get(ext, "mp3"))
            calls.append(OmniCall("audio", filename or "audio", latency_ms, usage, None))
            units = [Unit(None, "transcript", text, omni_used=True)]
        except omni_video.OmniError as err:
            calls.append(OmniCall("audio", filename or "audio", None, None, str(err)))
            units = [Unit(None, "transcript", "", omni_used=True)]
    elif kind == "video":
        try:
            text, usage, latency_ms = await omni_video.transcribe_video(data, mime=VIDEO_EXTS.get(ext, "video/mp4"))
            calls.append(OmniCall("video", filename or "video", latency_ms, usage, None))
            units = [Unit(None, "transcript", text, omni_used=True)]
        except omni_video.OmniError as err:
            calls.append(OmniCall("video", filename or "video", None, None, str(err)))
            units = [Unit(None, "transcript", "", omni_used=True)]
    else:
        units = [Unit(None, "text", data.decode("utf-8", errors="replace"), omni_used=False)]

    return ProcessResult(kind=kind, units=units, omni_calls=calls)


@router.post("/file-processing/process")
async def process_file_endpoint(file: UploadFile = File(...)):
    filename = file.filename or "upload"
    data = await file.read()
    if not data:
        raise HTTPException(status_code=400, detail="Empty file")

    # Omni-configured or not, process_file runs the same path either way: a page/image/audio/
    # video call that needs Omni without a configured key surfaces OmniNotConfigured as that
    # unit's own error (same fail-open rule as the video pipelines), rather than this endpoint
    # trying to pre-guess which kinds would have needed Omni and skip them itself.
    do_process = omni_video.is_configured()
    task_start = time.perf_counter()
    result = await process_file(data, filename=filename)
    latency_ms = round((time.perf_counter() - task_start) * 1000)

    omni_calls_info = [
        {"kind": c.kind, "label": c.label, "latency_ms": c.latency_ms, "tokens": c.tokens, "error": c.error}
        for c in result.omni_calls
    ]
    tokens_total = sum(c.tokens["total_tokens"] for c in result.omni_calls if c.tokens)
    units_with_omni = sum(1 for u in result.units if u.omni_used)
    omni_latencies = [c.latency_ms for c in result.omni_calls if c.latency_ms is not None]
    # Sum: total Omni compute across every call — NOT wall clock, since calls within a page/image
    # batch run concurrently (asyncio.gather). Max: the single slowest call, which is roughly the
    # wall-clock floor those concurrent calls actually cost, closer to what a person feels waiting.
    omni_latency_sum_ms = sum(omni_latencies)
    omni_latency_max_ms = max(omni_latencies) if omni_latencies else 0

    combined_text = "\n\n".join(
        f"=== {'Page/Unit ' + str(u.index) if u.index is not None else 'Document'} ({u.content_type}) ===\n{u.text}"
        for u in result.units
    )

    return {
        "filename": filename,
        "kind": result.kind,
        "unit_count": len(result.units),
        "units_using_omni": units_with_omni,
        "omni_configured": do_process,
        "omni_key_env": omni_video.OMNI_API_KEY_ENV,
        "omni_calls": omni_calls_info,
        "omni_call_count": len(omni_calls_info),
        "omni_tokens_total": tokens_total,
        "omni_latency_ms": omni_latency_sum_ms,
        "omni_latency_max_ms": omni_latency_max_ms,
        "latency_ms": latency_ms,
        "units": [
            {"index": u.index, "content_type": u.content_type, "text": u.text, "omni_used": u.omni_used}
            for u in result.units
        ],
        "combined_text": combined_text,
    }
